"""Agreement templates (Tamil Nadu practice, in English) and rendering to Word (.docx) and PDF in stamp-paper layout:
the first page keeps the top part blank for the non-judicial stamp paper / e-stamp, wide left binding margin, and every
page carries the parties' initials line and 'Page x of y'. Templates are drafts for a lawyer to review before use."""
import math
import re
from datetime import date
from io import BytesIO
from logic import amount_in_words

PAPER = {"legal": (215.9, 355.6), "a4": (210.0, 297.0)}  # mm

# Common blocks of fields
def _party(key, label):
    return [
        {"key": f"{key}_name", "label": f"{label}: name", "required": True, "group": label},
        {"key": f"{key}_relation", "label": f"{label}: S/o, D/o, W/o (or 'represented by')", "group": label},
        {"key": f"{key}_age", "label": f"{label}: age", "type": "number", "group": label},
        {"key": f"{key}_address", "label": f"{label}: address", "type": "textarea", "required": True, "group": label},
        {"key": f"{key}_id", "label": f"{label}: PAN / Aadhaar (last 4 only)", "group": label},
    ]


EXEC = [{"key": "execution_place", "label": "Place of signing", "required": True, "default": "Erode", "group": "Signing"},
        {"key": "execution_date", "label": "Date of signing", "type": "date", "required": True, "group": "Signing"}]
PROPERTY = [{"key": "property_address", "label": "Property address", "type": "textarea", "required": True, "group": "Property"},
            {"key": "property_description", "label": "Description (type, rooms, floor, fittings)", "type": "textarea", "group": "Property"},
            {"key": "survey_no", "label": "Survey / door / patta no.", "group": "Property"},
            {"key": "boundaries", "label": "Boundaries (North / South / East / West)", "type": "textarea", "group": "Property"}]
COMPANY = [{"key": "company_signatory", "label": "Company signatory", "required": True, "default": "", "group": "Company"},
           {"key": "company_signatory_designation", "label": "Signatory's designation", "default": "Director", "group": "Company"}]


def _lease_stamp(d):
    months = int(float(d.get("term_months") or 11))
    rent = float(d.get("rent") or 0)
    total_rent = 0.0
    esc, every = float(d.get("escalation_pct") or 0), int(float(d.get("escalation_every_months") or 12) or 12)
    for m in range(months):
        total_rent += rent * (1 + esc / 100) ** (m // every)
    adv = float(d.get("deposit") or 0) + float(d.get("advance_rent") or 0)
    base = total_rent + adv
    return {"basis": round(base, 2), "rate": "1% (Art. 35(a), leases under 30 years)",
            "duty": math.ceil(base * 0.01),
            "note": "Tamil Nadu: 1% of the total rent for the term plus the advance and security deposit (from 3 May 2024). "
                    "Registration fee is separate. Confirm with the sub-registrar before buying stamp paper."}


def _flat(amount, article):
    return lambda d: {"basis": None, "rate": article, "duty": amount,
                      "note": "Confirm the current stamp value with the sub-registrar or your lawyer."}


def _sale_stamp(d):
    if d.get("possession_on_agreement") in (True, "yes", "true"):
        price = float(d.get("price") or 0)
        return {"basis": price, "rate": "7% (agreement with possession is charged like a sale)", "duty": math.ceil(price * 0.07),
                "note": "When possession is handed over under the agreement it is charged like a sale deed. Confirm with your lawyer."}
    return {"basis": None, "rate": "Agreement (Art. 5)", "duty": 200,
            "note": "Agreement for sale without possession. Confirm the current stamp value and whether to register it."}


TEMPLATES = {
    "rental_residential": {
        "title": "RENTAL AGREEMENT", "subtitle": "Residential premises", "group": "Rent & lease",
        "summary": "11-month house / flat rental between owner and tenant.",
        "parties": [("lessor", "LESSOR (Owner)"), ("lessee", "LESSEE (Tenant)")],
        "sources": ["lease", "property"],
        "fields": EXEC + _party("lessor", "Lessor") + _party("lessee", "Lessee") + PROPERTY + [
            {"key": "rent", "label": "Monthly rent (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "deposit", "label": "Security deposit (Rs.)", "type": "money", "group": "Terms"},
            {"key": "start_date", "label": "Tenancy starts", "type": "date", "required": True, "group": "Terms"},
            {"key": "term_months", "label": "Term (months)", "type": "number", "default": 11, "group": "Terms"},
            {"key": "due_day", "label": "Rent due by day of month", "type": "number", "default": 5, "group": "Terms"},
            {"key": "escalation_pct", "label": "Increase on renewal (%)", "type": "number", "default": 5, "group": "Terms"},
            {"key": "notice_months", "label": "Notice period (months)", "type": "number", "default": 1, "group": "Terms"},
            {"key": "lock_in_months", "label": "Lock-in (months, 0 for none)", "type": "number", "default": 0, "group": "Terms"},
            {"key": "maintenance_by", "label": "Apartment maintenance paid by", "type": "select", "options": ["Lessee", "Lessor"], "default": "Lessee", "group": "Terms"},
            {"key": "fittings", "label": "Fittings and furniture provided", "type": "textarea", "group": "Terms"},
            {"key": "use", "label": "Use", "default": "residential purpose for the Lessee and family only", "group": "Terms"},
        ],
        "recitals": ["This RENTAL AGREEMENT is made and executed at {execution_place} on {execution_date} BETWEEN:",
                     "{lessor_block}, hereinafter called the \"LESSOR\" (which expression shall include heirs, legal representatives and assigns) of the ONE PART;",
                     "AND",
                     "{lessee_block}, hereinafter called the \"LESSEE\" (which expression shall include heirs, legal representatives and permitted assigns) of the OTHER PART.",
                     "WHEREAS the Lessor is the absolute owner of the premises described in the Schedule below (the \"Premises\"), and the Lessee has requested the Lessor to let out the Premises, and the Lessor has agreed to do so on the following terms:",
                     "NOW THIS AGREEMENT WITNESSES AS FOLLOWS:"],
        "clauses": [
            "TERM: The tenancy is for a period of {term_months} months commencing from {start_date} and ending on {end_date}, and may be renewed by mutual written consent.",
            "RENT: The Lessee shall pay a monthly rent of {rent_full} on or before the {due_day_th} day of every English calendar month, by bank transfer / UPI to the Lessor's account or as the Lessor directs in writing. The Lessor shall give a receipt for every payment.",
            "SECURITY DEPOSIT: The Lessee has paid {deposit_full} as an interest-free refundable security deposit, the receipt of which the Lessor acknowledges. It shall be refunded when the Lessee hands over vacant possession, after deducting unpaid rent, unpaid charges and the cost of repairing damage beyond normal wear and tear.",
            "RENEWAL: On renewal the rent shall be increased by {escalation_pct}% unless the parties agree otherwise in writing.",
            "CHARGES: The Lessee shall pay electricity charges as per the separate meter and water charges, if any, directly to the authorities. {maintenance_by} shall pay the apartment / association maintenance charges. Property tax shall be paid by the Lessor.",
            "USE: The Premises shall be used only for {use} and not for any illegal, immoral or commercial purpose. The Lessee shall not sub-let, assign or part with possession of the Premises or any part of it.",
            "CARE OF PREMISES: The Lessee shall keep the Premises and the fittings listed in the Schedule in good and tenantable condition and shall not make structural alterations without the Lessor's written consent. Minor day-to-day repairs are the Lessee's responsibility; major and structural repairs are the Lessor's.",
            "INSPECTION: The Lessor or the Lessor's authorised agent may inspect the Premises at a reasonable time after giving prior notice to the Lessee.",
            "LOCK-IN AND NOTICE: {lock_in_text}Either party may end this tenancy by giving {notice_months} month(s) notice in writing to the other.",
            "DEFAULT: If the Lessee fails to pay the rent for two consecutive months or breaches any term of this agreement, the Lessor may terminate this agreement after giving 15 days written notice, and the Lessee shall hand over vacant possession.",
            "HANDOVER: On expiry or termination the Lessee shall hand over vacant and peaceful possession of the Premises with all fittings, together with the keys, and the final meter readings shall be recorded.",
            "TENANCY LAW: The parties shall comply with the Tamil Nadu Regulation of Rights and Responsibilities of Landlords and Tenants Act, 2017, and the Lessor shall inform the Rent Authority of this tenancy as that Act requires.",
            "This agreement is made in duplicate; the original is kept by the Lessor and the duplicate by the Lessee.",
        ],
        "schedule": "SCHEDULE OF PREMISES\n{property_address}\n{property_description}\n{survey_line}{boundaries_line}{fittings_line}",
        "stamp": _lease_stamp,
    },
    "lease_commercial": {
        "title": "LEASE DEED", "subtitle": "Commercial / industrial premises", "group": "Rent & lease",
        "summary": "Shop, office, warehouse or factory shed lease with lock-in and escalation.",
        "parties": [("lessor", "LESSOR"), ("lessee", "LESSEE")],
        "sources": ["lease", "property", "deal"],
        "fields": EXEC + _party("lessor", "Lessor") + _party("lessee", "Lessee") + PROPERTY + [
            {"key": "rent", "label": "Monthly rent (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "deposit", "label": "Security deposit (Rs.)", "type": "money", "group": "Terms"},
            {"key": "start_date", "label": "Lease starts", "type": "date", "required": True, "group": "Terms"},
            {"key": "term_months", "label": "Term (months)", "type": "number", "default": 36, "group": "Terms"},
            {"key": "lock_in_months", "label": "Lock-in (months)", "type": "number", "default": 12, "group": "Terms"},
            {"key": "escalation_pct", "label": "Escalation (%)", "type": "number", "default": 5, "group": "Terms"},
            {"key": "escalation_every_months", "label": "Escalation every (months)", "type": "number", "default": 12, "group": "Terms"},
            {"key": "rent_free_days", "label": "Rent-free fit-out period (days)", "type": "number", "default": 0, "group": "Terms"},
            {"key": "due_day", "label": "Rent due by day of month", "type": "number", "default": 5, "group": "Terms"},
            {"key": "notice_months", "label": "Notice after lock-in (months)", "type": "number", "default": 3, "group": "Terms"},
            {"key": "use", "label": "Permitted use", "default": "the Lessee's lawful business", "group": "Terms"},
            {"key": "power", "label": "Power sanctioned (HP / kVA)", "group": "Terms"},
            {"key": "gst_note", "label": "GST on rent", "type": "select", "options": ["GST extra as applicable", "Rent is inclusive of GST", "Not applicable"], "default": "GST extra as applicable", "group": "Terms"},
        ],
        "recitals": ["This LEASE DEED is made and executed at {execution_place} on {execution_date} BETWEEN:",
                     "{lessor_block}, hereinafter called the \"LESSOR\" of the ONE PART;", "AND",
                     "{lessee_block}, hereinafter called the \"LESSEE\" of the OTHER PART.",
                     "WHEREAS the Lessor is the absolute owner of the property described in the Schedule (the \"Premises\") and has agreed to lease it to the Lessee for {use} on the terms below.",
                     "NOW THIS DEED WITNESSES AS FOLLOWS:"],
        "clauses": [
            "TERM: The lease is for {term_months} months from {start_date} to {end_date}, renewable on terms agreed in writing.",
            "FIT-OUT: The Lessee is allowed a rent-free period of {rent_free_days} days from the start date for interior works at its own cost.",
            "RENT: The Lessee shall pay a monthly rent of {rent_full}, {gst_note}, on or before the {due_day_th} of each month. Rent is payable after deducting tax at source as required by law, and the Lessee shall issue TDS certificates.",
            "ESCALATION: The rent shall increase by {escalation_pct}% every {escalation_every_months} months on the then prevailing rent.",
            "SECURITY DEPOSIT: The Lessee has paid {deposit_full} as an interest-free refundable deposit, refundable on handover of vacant possession after adjusting dues and damages.",
            "LOCK-IN: Neither party shall terminate this lease during the first {lock_in_months} months, except for breach. If the Lessee leaves during the lock-in, rent for the balance lock-in period shall be payable.",
            "TERMINATION: After the lock-in, either party may terminate by {notice_months} months written notice.",
            "USE AND COMPLIANCE: The Premises shall be used only for {use}. The Lessee shall obtain and keep its own trade licences, pollution and fire approvals for its business and comply with all laws.",
            "UTILITIES: Electricity (sanctioned load {power}), water and other consumption charges shall be paid by the Lessee. Property tax shall be paid by the Lessor.",
            "REPAIRS: Structural repairs, roof and external walls are the Lessor's responsibility; day-to-day maintenance and repairs to the Lessee's installations are the Lessee's.",
            "ALTERATIONS: The Lessee may install partitions, machinery and fixtures that can be removed without damage; structural changes need the Lessor's written consent. On handover the Lessee shall restore the Premises to their original condition except normal wear and tear.",
            "INSURANCE: The Lessor shall insure the building; the Lessee shall insure its stock, machinery and fit-outs.",
            "ASSIGNMENT: The Lessee shall not assign or sub-let the Premises without the Lessor's written consent.",
            "INDEMNITY: Each party shall indemnify the other against losses arising from its own breach of this deed.",
            "DISPUTES: Disputes shall first be resolved by discussion, failing which the courts at {execution_place} shall have jurisdiction.",
        ],
        "schedule": "SCHEDULE OF PREMISES\n{property_address}\n{property_description}\n{survey_line}{boundaries_line}",
        "stamp": _lease_stamp,
    },
    "leave_licence": {
        "title": "LEAVE AND LICENCE AGREEMENT", "subtitle": "", "group": "Rent & lease",
        "summary": "Licence to occupy (no tenancy created), common for corporate stays and short terms.",
        "parties": [("lessor", "LICENSOR"), ("lessee", "LICENSEE")],
        "sources": ["lease", "property"],
        "fields": EXEC + _party("lessor", "Licensor") + _party("lessee", "Licensee") + PROPERTY + [
            {"key": "rent", "label": "Monthly licence fee (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "deposit", "label": "Security deposit (Rs.)", "type": "money", "group": "Terms"},
            {"key": "start_date", "label": "Licence starts", "type": "date", "required": True, "group": "Terms"},
            {"key": "term_months", "label": "Term (months)", "type": "number", "default": 11, "group": "Terms"},
            {"key": "due_day", "label": "Fee due by day of month", "type": "number", "default": 5, "group": "Terms"},
            {"key": "notice_months", "label": "Notice (months)", "type": "number", "default": 1, "group": "Terms"},
        ],
        "recitals": ["This LEAVE AND LICENCE AGREEMENT is made at {execution_place} on {execution_date} BETWEEN:",
                     "{lessor_block}, hereinafter called the \"LICENSOR\";", "AND", "{lessee_block}, hereinafter called the \"LICENSEE\".",
                     "The Licensor permits the Licensee to use the premises described in the Schedule on leave and licence basis on the following terms:"],
        "clauses": [
            "The licence is for {term_months} months from {start_date} to {end_date}.",
            "The Licensee shall pay a monthly licence fee of {rent_full} on or before the {due_day_th} of each month.",
            "The Licensee has paid {deposit_full} as an interest-free refundable deposit, refundable on vacating after adjusting dues.",
            "This agreement is a licence only. It does not create any tenancy, lease or other right or interest in the premises in favour of the Licensee, and legal possession remains with the Licensor.",
            "The Licensee shall use the premises only for residence and shall not transfer, sub-license or permit any other person to occupy them.",
            "The Licensee shall pay electricity and other consumption charges and keep the premises in good condition.",
            "Either party may end the licence by giving {notice_months} month(s) notice in writing. On expiry or termination the Licensee shall vacate and hand over the premises with all fittings.",
        ],
        "schedule": "SCHEDULE OF PREMISES\n{property_address}\n{property_description}\n{survey_line}",
        "stamp": _lease_stamp,
    },
    "sale_agreement": {
        "title": "AGREEMENT FOR SALE", "subtitle": "", "group": "Sale",
        "summary": "Seller and buyer agree the price, advance, balance date and title conditions before the sale deed.",
        "parties": [("seller", "VENDOR (Seller)"), ("buyer", "PURCHASER (Buyer)")],
        "sources": ["deal", "property"],
        "fields": EXEC + _party("seller", "Vendor") + _party("buyer", "Purchaser") + PROPERTY + [
            {"key": "price", "label": "Total sale price (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "advance", "label": "Advance paid today (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "advance_mode", "label": "Advance paid by", "default": "bank transfer / cheque no. ________", "group": "Terms"},
            {"key": "completion_days", "label": "Balance to be paid within (days)", "type": "number", "default": 90, "group": "Terms"},
            {"key": "possession_on_agreement", "label": "Possession handed over now?", "type": "select", "options": ["no", "yes"], "default": "no", "group": "Terms"},
            {"key": "forfeit_note", "label": "If the buyer fails", "default": "the Vendor may forfeit 10% of the sale price from the advance and refund the rest", "group": "Terms"},
        ],
        "recitals": ["This AGREEMENT FOR SALE is made at {execution_place} on {execution_date} BETWEEN:",
                     "{seller_block}, hereinafter called the \"VENDOR\" of the ONE PART;", "AND", "{buyer_block}, hereinafter called the \"PURCHASER\" of the OTHER PART.",
                     "WHEREAS the Vendor is the absolute owner in possession of the property described in the Schedule, having clear and marketable title, and has agreed to sell it to the Purchaser, who has agreed to buy it, on the terms below:"],
        "clauses": [
            "PRICE: The total sale consideration is {price_full}.",
            "ADVANCE: The Purchaser has today paid {advance_full} by {advance_mode}, which the Vendor acknowledges. The balance of {balance_full} shall be paid at the time of registration of the sale deed.",
            "COMPLETION: The sale shall be completed and the sale deed registered within {completion_days} days from today, that is on or before {completion_date}.",
            "TITLE: The Vendor shall hand over copies of the parent documents, patta, encumbrance certificate for 30 years, tax receipts and approvals for verification by the Purchaser's lawyer within 7 days. If the title is found defective and is not cured, the Vendor shall refund the advance in full.",
            "ENCUMBRANCES: The Vendor declares that the property is free from mortgages, attachments, litigation and claims, and shall not deal with it in any manner until the sale is completed.",
            "TAXES: The Vendor shall pay all property tax, electricity and other dues up to the date of registration.",
            "COSTS: Stamp duty, registration charges and documentation costs of the sale deed shall be borne by the Purchaser. If the sale value is Rs. 50 lakh or more, the Purchaser shall deduct and deposit tax at source (1%) and give the Vendor the certificate.",
            "DEFAULT BY PURCHASER: If the Purchaser fails to pay the balance within the agreed time, {forfeit_note}.",
            "DEFAULT BY VENDOR: If the Vendor fails to complete the sale, the Purchaser may seek specific performance of this agreement or refund of the advance with an equal amount as compensation.",
            "POSSESSION: Vacant possession shall be handed over on registration of the sale deed{possession_text}.",
        ],
        "schedule": "SCHEDULE OF PROPERTY\n{property_address}\n{property_description}\n{survey_line}{boundaries_line}",
        "stamp": _sale_stamp,
    },
    "token_receipt": {
        "title": "TOKEN ADVANCE RECEIPT AND BOOKING MEMORANDUM", "subtitle": "", "group": "Sale",
        "summary": "Receipt for a token advance that holds a property while documents are checked.",
        "parties": [("seller", "OWNER"), ("buyer", "BUYER / TENANT")],
        "sources": ["deal", "property"],
        "fields": EXEC + _party("seller", "Owner") + _party("buyer", "Buyer") + PROPERTY + [
            {"key": "price", "label": "Agreed price / rent (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "advance", "label": "Token amount (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "advance_mode", "label": "Paid by", "default": "UPI / bank transfer", "group": "Terms"},
            {"key": "hold_days", "label": "Property held for (days)", "type": "number", "default": 15, "group": "Terms"},
        ],
        "recitals": ["Received at {execution_place} on {execution_date} by {seller_name} (the \"Owner\") from {buyer_name} (the \"Buyer\") the sum of {advance_full} by {advance_mode} as token advance for the property described below, on the following terms:"],
        "clauses": [
            "The agreed price is {price_full}. The token is part of this price.",
            "The Owner shall not deal with the property with anyone else for {hold_days} days from today, during which the Buyer's lawyer will verify the documents.",
            "A detailed agreement will be signed within {hold_days} days. If the title is not clear, the token shall be refunded in full within 7 days.",
            "If the Buyer withdraws without a valid reason, the token may be retained by the Owner.",
        ],
        "schedule": "PROPERTY\n{property_address}\n{property_description}\n{survey_line}",
        "stamp": _flat(200, "Agreement / receipt (confirm)"),
    },
    "brokerage_authorisation": {
        "title": "MARKETING AUTHORISATION AND BROKERAGE AGREEMENT", "subtitle": "", "group": "Owner services",
        "summary": "Owner authorises the company to market the property and agrees the brokerage.",
        "parties": [("owner", "OWNER"), ("company", "SELLANDIAMMAN ESTATES")],
        "sources": ["property", "owner"],
        "fields": EXEC + _party("owner", "Owner") + COMPANY + PROPERTY + [
            {"key": "transaction", "label": "Owner wants to", "type": "select", "options": ["sell", "rent out", "lease"], "default": "sell", "group": "Terms"},
            {"key": "expected_price", "label": "Expected price / rent (Rs.)", "type": "money", "group": "Terms"},
            {"key": "brokerage", "label": "Brokerage", "default": "2% of the sale price", "group": "Terms"},
            {"key": "exclusive", "label": "Exclusive?", "type": "select", "options": ["exclusive", "non-exclusive"], "default": "non-exclusive", "group": "Terms"},
            {"key": "valid_months", "label": "Valid for (months)", "type": "number", "default": 6, "group": "Terms"},
        ],
        "recitals": ["This agreement is made at {execution_place} on {execution_date} between {owner_block} (the \"Owner\") and {company_block} (the \"Company\")."],
        "clauses": [
            "AUTHORISATION: The Owner authorises the Company, on a {exclusive} basis, to market the property described in the Schedule to {transaction} it, at an expected value of {expected_price_full}, on its website, partner portals and social media, and to show it to prospective customers.",
            "PERIOD: This authorisation is valid for {valid_months} months from today and may be extended in writing.",
            "BROKERAGE: The Owner shall pay the Company brokerage of {brokerage} (plus GST if applicable) on completion of a transaction with any customer introduced by the Company, including transactions completed within 6 months after this agreement ends with such a customer.",
            "OWNER'S DECLARATION: The Owner declares that the Owner has clear title and full authority to deal with the property and that the information given to the Company is true.",
            "PHOTOGRAPHS: The Owner permits the Company to photograph the property and use the photographs for marketing during this agreement. The exact address will be shared only with verified enquiries.",
            "NO GUARANTEE: The Company will make reasonable efforts but does not guarantee a transaction or a price. The Owner is free to accept or reject any offer.",
            "The Company's TNRERA agent registration number, where applicable, is stated in its letterhead.",
        ],
        "schedule": "PROPERTY\n{property_address}\n{property_description}\n{survey_line}",
        "stamp": _flat(200, "Agreement not otherwise provided, Art. 5(j)"),
    },
    "property_management": {
        "title": "PROPERTY MANAGEMENT AGREEMENT", "subtitle": "", "group": "Owner services",
        "summary": "Company finds tenants, collects rent, arranges repairs and sends monthly statements.",
        "parties": [("owner", "OWNER"), ("company", "SELLANDIAMMAN ESTATES")],
        "sources": ["property", "owner"],
        "fields": EXEC + _party("owner", "Owner") + COMPANY + PROPERTY + [
            {"key": "mgmt_fee_pct", "label": "Management fee (% of rent collected)", "type": "number", "default": 8, "group": "Terms"},
            {"key": "letting_fee", "label": "Tenant-finding fee", "default": "one month's rent per new tenancy", "group": "Terms"},
            {"key": "repair_limit", "label": "Repairs without asking up to (Rs.)", "type": "money", "default": 5000, "group": "Terms"},
            {"key": "statement_day", "label": "Monthly statement by day", "type": "number", "default": 10, "group": "Terms"},
            {"key": "term_months", "label": "Term (months)", "type": "number", "default": 12, "group": "Terms"},
            {"key": "notice_months", "label": "Notice (months)", "type": "number", "default": 2, "group": "Terms"},
            {"key": "owner_bank", "label": "Owner's bank account (bank, last 4 digits)", "group": "Terms"},
        ],
        "recitals": ["This agreement is made at {execution_place} on {execution_date} between {owner_block} (the \"Owner\") and {company_block} (the \"Company\").",
                     "The Owner appoints the Company to manage the property described in the Schedule on the following terms:"],
        "clauses": [
            "SERVICES: The Company shall advertise the property, screen and select tenants with the Owner's approval, prepare rental agreements, hand over and take back the property with inspection reports, collect rent and deposits, follow up arrears, and arrange repairs and maintenance.",
            "OWNER FUNDS: Rent and deposits are collected on behalf of the Owner and held in a separate owner-funds account. They are never the Company's income.",
            "FEES: The Company's fee is {mgmt_fee_pct}% of the rent actually collected, plus {letting_fee}, plus GST if applicable. Fees are deducted in the monthly statement.",
            "REPAIRS: The Company may arrange repairs up to {repair_limit_full} per job without prior approval; larger work needs the Owner's written approval and a quotation. Costs are recovered in the next statement.",
            "STATEMENTS AND PAYOUTS: The Company shall send a statement by the {statement_day_th} of each month showing rent collected, fees, expenses and the net amount, and shall transfer the net amount to the Owner's account {owner_bank} with the bank reference recorded on the statement. The Owner can also see these in the owner portal.",
            "TERM: This agreement is for {term_months} months and continues thereafter until ended by either party giving {notice_months} months notice in writing. Tenancies already in force continue on their terms.",
            "OWNER'S DUTIES: The Owner shall keep the property insured, pay property tax and keep title documents in order.",
            "LIABILITY: The Company acts as the Owner's agent and is not liable for a tenant's default or damage caused by a tenant, but shall take all reasonable steps to recover them.",
        ],
        "schedule": "PROPERTY\n{property_address}\n{property_description}\n{survey_line}",
        "stamp": _flat(200, "Agreement not otherwise provided, Art. 5(j)"),
    },
    "shortstay_management": {
        "title": "SHORT-STAY MANAGEMENT AGREEMENT", "subtitle": "Holiday home / serviced stay", "group": "Owner services",
        "summary": "Company runs the home on Airbnb, Booking.com and other platforms and settles each stay with the owner.",
        "parties": [("owner", "OWNER"), ("company", "SELLANDIAMMAN ESTATES")],
        "sources": ["property", "owner"],
        "fields": EXEC + _party("owner", "Owner") + COMPANY + PROPERTY + [
            {"key": "mgmt_fee_pct", "label": "Management fee (% of booking value after platform fees)", "type": "number", "default": 20, "group": "Terms"},
            {"key": "platforms", "label": "Platforms", "default": "Airbnb, Booking.com, MakeMyTrip, Agoda and direct bookings", "group": "Terms"},
            {"key": "owner_nights", "label": "Owner's own nights per year", "type": "number", "default": 30, "group": "Terms"},
            {"key": "term_months", "label": "Term (months)", "type": "number", "default": 12, "group": "Terms"},
            {"key": "notice_months", "label": "Notice (months)", "type": "number", "default": 2, "group": "Terms"},
        ],
        "recitals": ["This agreement is made at {execution_place} on {execution_date} between {owner_block} (the \"Owner\") and {company_block} (the \"Company\").",
                     "The Owner appoints the Company to operate the furnished home described in the Schedule for short stays on the following terms:"],
        "clauses": [
            "LISTINGS: The Company shall list and manage the home on {platforms}, keep one shared calendar so that dates are not double-booked, set prices, and communicate with guests.",
            "OPERATIONS: The Company shall arrange check-in, housekeeping and linen, turnover cleaning, supplies and guest support, and keep a record of guest issues.",
            "GUEST RECORDS: The Company shall collect guest identity and submit foreign guests' Form C as required by law.",
            "FEES: The Company's fee is {mgmt_fee_pct}% of the booking value remaining after platform commissions. Operating costs (cleaning, supplies, utilities during stays) are deducted at actuals.",
            "SETTLEMENT: Each completed stay is settled with a statement showing gross booking value, platform fee, management fee, costs and the owner's payout, which is transferred to the Owner's account and visible in the owner portal.",
            "OWNER STAYS: The Owner may block up to {owner_nights} nights a year for personal use, with 14 days notice. Confirmed guest bookings cannot be cancelled for owner stays.",
            "LICENCES AND TAXES: The Owner shall obtain any licence required by the local body and pay property tax; GST and income tax on the Owner's earnings are the Owner's responsibility.",
            "DAMAGE: Guest damage shall be recovered from the guest or the platform's host protection where possible.",
            "TERM: {term_months} months, continuing until ended by {notice_months} months notice; bookings already confirmed shall be honoured.",
        ],
        "schedule": "PROPERTY\n{property_address}\n{property_description}",
        "stamp": _flat(200, "Agreement not otherwise provided, Art. 5(j)"),
    },
    "investor_coinvestment": {
        "title": "CO-INVESTMENT AGREEMENT", "subtitle": "", "group": "Investors",
        "summary": "Investor's contribution to a project, share, payouts, reporting and exit.",
        "parties": [("investor", "INVESTOR"), ("company", "SELLANDIAMMAN ESTATES")],
        "sources": ["investment", "investor"],
        "fields": EXEC + _party("investor", "Investor") + COMPANY + [
            {"key": "project_name", "label": "Project", "required": True, "group": "Project"},
            {"key": "project_location", "label": "Project location", "group": "Project"},
            {"key": "project_description", "label": "Project description", "type": "textarea", "group": "Project"},
            {"key": "amount", "label": "Investment amount (Rs.)", "type": "money", "required": True, "group": "Terms"},
            {"key": "investment_date", "label": "Date money received", "type": "date", "group": "Terms"},
            {"key": "reference", "label": "Payment reference", "group": "Terms"},
            {"key": "horizon_months", "label": "Expected holding (months)", "type": "number", "default": 36, "group": "Terms"},
            {"key": "payout_frequency", "label": "Payouts", "default": "quarterly, from net rental income", "group": "Terms"},
            {"key": "mgmt_fee", "label": "Company's management fee", "default": "10% of net income", "group": "Terms"},
        ],
        "recitals": ["This agreement is made at {execution_place} on {execution_date} between {investor_block} (the \"Investor\") and {company_block} (the \"Company\").",
                     "The Company has organised the project \"{project_name}\" at {project_location} ({project_description}) and the Investor wishes to participate on the following terms:"],
        "clauses": [
            "CONTRIBUTION: The Investor has contributed {amount_full} on {investment_date} (reference {reference}), which the Company acknowledges.",
            "SHARE: The Investor's share of the project's income and value is in proportion to the Investor's contribution to the total capital contributed by all investors, as shown in the investor dashboard and statements.",
            "PAYOUTS: Income shall be distributed {payout_frequency}, after expenses and the Company's management fee of {mgmt_fee}, with tax deducted at source where the law requires.",
            "REPORTING: The Company shall give the Investor access to an online dashboard and statements showing the contribution, current estimated value, payouts and returns, and shall post project updates.",
            "VALUATION: Current values are the Company's estimates based on valuations recorded from time to time and are not offers or guarantees.",
            "HOLDING AND EXIT: The expected holding period is {horizon_months} months. Exit before that is subject to the Company finding a buyer for the Investor's share or another arrangement agreed in writing.",
            "RISK: The Investor understands that real estate values and income can fall as well as rise, that returns are not guaranteed, and that the investment is not a deposit.",
            "CONFIDENTIALITY AND DISPUTES: Both parties shall keep the terms confidential. Disputes shall be resolved by discussion, failing which by arbitration at {execution_place} under the Arbitration and Conciliation Act, 1996.",
        ],
        "schedule": "",
        "stamp": _flat(200, "Agreement not otherwise provided, Art. 5(j)"),
    },
}


def template_list():
    return [{"key": k, "title": t["title"].title(), "subtitle": t["subtitle"], "group": t["group"], "summary": t["summary"],
             "parties": [p[1] for p in t["parties"]], "sources": t["sources"], "fields": t["fields"]} for k, t in TEMPLATES.items()]


# ---------------- Rendering ----------------
def money_full(v):
    v = float(v or 0)
    # Indian digit grouping: 12,34,567
    whole, frac = f"{v:.2f}".split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        whole = ",".join(parts) + "," + tail
    s = whole + ("" if frac == "00" else f".{frac}")
    return f"Rs. {s}/- ({amount_in_words(v)})"


def nice_date(s):
    try:
        d = date.fromisoformat(str(s)[:10])
        return f"{d.day} {d.strftime('%B %Y')}"
    except (ValueError, TypeError):
        return str(s or "________")


def _th(n):
    try:
        n = int(float(n))
    except (TypeError, ValueError):
        return str(n or "")
    return f"{n}{'th' if 11 <= n % 100 <= 13 else {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')}"


def _add_months(iso, months, minus_day=True):
    from datetime import timedelta
    d = date.fromisoformat(iso)
    m = d.month - 1 + int(months)
    y, m = d.year + m // 12, m % 12 + 1
    import calendar
    day = min(d.day, calendar.monthrange(y, m)[1])
    out = date(y, m, day)
    return (out - timedelta(days=1)) if minus_day else out


def context(tpl_key, data, company):
    t = TEMPLATES[tpl_key]
    d = {f["key"]: f.get("default") for f in t["fields"]}
    d.update({k: v for k, v in (data or {}).items() if v not in (None, "")})
    blank = "________________"
    ctx = {k: (blank if v in (None, "") else v) for k, v in d.items()}

    def block(key, label):
        if key == "company":
            return (f"{company.get('legal_name') or company.get('name') or 'Sellandiamman Estates'}, having its office at "
                    f"{company.get('address') or blank}, represented by its {d.get('company_signatory_designation') or 'authorised signatory'} "
                    f"{d.get('company_signatory') or blank}")
        parts = [str(d.get(f"{key}_name") or blank)]
        if d.get(f"{key}_relation"):
            parts.append(str(d[f"{key}_relation"]))
        if d.get(f"{key}_age"):
            parts.append(f"aged about {int(float(d[f'{key}_age']))} years")
        business = re.search(r"(ltd|llp|pvt|traders|company|industries|logistics|textiles|enterprises)", str(d.get(f"{key}_name") or "").lower())
        parts.append(f"{'having its office at' if business else 'residing at'} {d.get(f'{key}_address') or blank}")
        if d.get(f"{key}_id"):
            parts.append(f"(ID ending {str(d[f'{key}_id'])[-4:]})")
        return ", ".join(parts)
    for key, _label in t["parties"]:
        ctx[f"{key}_block"] = block(key, _label)
    for f in t["fields"]:
        if f.get("type") == "money":
            ctx[f"{f['key']}_full"] = money_full(d.get(f["key"])) if d.get(f["key"]) not in (None, "") else blank
        if f.get("type") == "date":
            ctx[f["key"]] = nice_date(d.get(f["key"])) if d.get(f["key"]) else blank
    for k in ("due_day", "statement_day"):
        if k in d:
            ctx[f"{k}_th"] = _th(d.get(k))
    if d.get("start_date") and d.get("term_months"):
        try:
            ctx["end_date"] = nice_date(_add_months(str(d["start_date"]), int(float(d["term_months"]))).isoformat())
        except (ValueError, TypeError):
            ctx["end_date"] = blank
    else:
        ctx["end_date"] = blank
    if d.get("price") and d.get("advance"):
        ctx["balance_full"] = money_full(float(d["price"]) - float(d["advance"]))
    else:
        ctx["balance_full"] = blank
    if d.get("execution_date") and d.get("completion_days"):
        from datetime import timedelta
        try:
            ctx["completion_date"] = nice_date((date.fromisoformat(str(d["execution_date"])) + timedelta(days=int(float(d["completion_days"])))).isoformat())
        except (ValueError, TypeError):
            ctx["completion_date"] = blank
    else:
        ctx["completion_date"] = blank
    lock = int(float(d.get("lock_in_months") or 0)) if str(d.get("lock_in_months") or "0").replace(".", "").isdigit() else 0
    ctx["lock_in_text"] = f"Neither party shall end this tenancy during the first {lock} months. " if lock else ""
    ctx["possession_text"] = " (possession has been handed over today on this agreement)" if d.get("possession_on_agreement") == "yes" else ""
    ctx["survey_line"] = f"Survey / door no.: {d['survey_no']}\n" if d.get("survey_no") else ""
    ctx["boundaries_line"] = f"Boundaries: {d['boundaries']}\n" if d.get("boundaries") else ""
    ctx["fittings_line"] = f"Fittings provided: {d['fittings']}\n" if d.get("fittings") else ""
    ctx["expected_price_full"] = money_full(d["expected_price"]) if d.get("expected_price") else "the price agreed with the Owner"
    ctx["owner_bank"] = d.get("owner_bank") or "as notified by the Owner in writing"
    return ctx, d


class _Safe(dict):
    def __missing__(self, k):
        return "________________"


def fill(text, ctx):
    return text.format_map(_Safe({k: (str(v) if not isinstance(v, str) else v) for k, v in ctx.items()}))


def render_blocks(tpl_key, data, company, clauses=None, extra_clauses=None):
    """Returns the agreement as a list of blocks shared by the on-screen preview, Word and PDF output."""
    t = TEMPLATES[tpl_key]
    ctx, d = context(tpl_key, data, company)
    blocks = [{"t": "title", "text": t["title"]}]
    if t.get("subtitle"):
        blocks.append({"t": "subtitle", "text": t["subtitle"]})
    for r in t["recitals"]:
        txt = fill(r, ctx)
        blocks.append({"t": "center" if txt.strip() in ("AND", "BETWEEN:") else "para", "text": txt})
    body = clauses if clauses is not None else [fill(c, ctx) for c in t["clauses"]]
    body = [c for c in body if str(c).strip()] + [fill(c, ctx) for c in (extra_clauses or []) if str(c).strip()]
    for i, c in enumerate(body, 1):
        blocks.append({"t": "clause", "n": i, "text": c})
    if t.get("schedule"):
        sched = [ln for ln in fill(t["schedule"], ctx).split("\n") if ln.strip()]
        blocks.append({"t": "heading", "text": sched[0]})
        for ln in sched[1:]:
            blocks.append({"t": "para", "text": ln})
    blocks.append({"t": "para", "text": "IN WITNESS WHEREOF the parties have signed this document on the day, month and year first written above, of their own free will and in the presence of the witnesses below."})
    sigs = []
    for key, label in t["parties"]:
        name = (company.get("legal_name") or company.get("name")) if key == "company" else d.get(f"{key}_name")
        sub = f"by {d.get('company_signatory') or ''} ({d.get('company_signatory_designation') or ''})" if key == "company" else ""
        sigs.append({"label": label, "name": name or "", "sub": sub})
    blocks.append({"t": "signatures", "parties": sigs})
    blocks.append({"t": "witnesses"})
    return blocks, ctx, d


def default_clauses(tpl_key, data, company):
    ctx, _ = context(tpl_key, data, company)
    return [fill(c, ctx) for c in TEMPLATES[tpl_key]["clauses"]]


def initials_line(tpl_key):
    return "      ".join(f"{label.split(' (')[0].title()}: ____________" for _, label in TEMPLATES[tpl_key]["parties"])


# ---------------- Word (.docx) ----------------
def to_docx(blocks, tpl_key, paper="legal", stamp_top_mm=100, stamp_line=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    doc = Document()
    w, h = PAPER.get(paper, PAPER["legal"])
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(w), Mm(h)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(32), Mm(22), Mm(22), Mm(28)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.3

    def field(run, code):
        for tag, text in (("begin", None), (None, code), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = fp.add_run(initials_line(tpl_key) + "        Page ")
    r.font.size = Pt(9)
    r2 = fp.add_run()
    r2.font.size = Pt(9)
    field(r2, "PAGE")
    r3 = fp.add_run(" of ")
    r3.font.size = Pt(9)
    r4 = fp.add_run()
    r4.font.size = Pt(9)
    field(r4, "NUMPAGES")
    # blank top of the first page, for the stamp paper's printed stamp / e-stamp
    first = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    first.paragraph_format.space_before = Mm(max(0, float(stamp_top_mm or 0)))
    if stamp_line:
        first.add_run(stamp_line).font.size = Pt(9)
    else:
        first.paragraph_format.space_after = Pt(0)
    for b in blocks:
        if b["t"] in ("title", "subtitle", "heading", "center"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(b["text"])
            run.bold = b["t"] != "subtitle"
            run.underline = b["t"] in ("title", "heading")
            run.font.size = Pt(15 if b["t"] == "title" else 12)
            if b["t"] == "heading":
                p.paragraph_format.space_before = Pt(12)
        elif b["t"] == "para":
            p = doc.add_paragraph(b["text"])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif b["t"] == "clause":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Mm(9)
            p.paragraph_format.first_line_indent = Mm(-9)
            m = re.match(r"^([A-Z][A-Z ,'&/-]{2,40}):\s*(.*)$", b["text"], re.S)
            p.add_run(f"{b['n']}.\t")
            if m:
                p.add_run(m.group(1) + ": ").bold = True
                p.add_run(m.group(2))
            else:
                p.add_run(b["text"])
        elif b["t"] == "signatures":
            doc.add_paragraph()
            tbl = doc.add_table(rows=4, cols=len(b["parties"]))
            for i, s in enumerate(b["parties"]):
                tbl.cell(0, i).text = ""
                tbl.cell(1, i).text = "\n\n______________________________"
                tbl.cell(2, i).text = s["label"]
                tbl.cell(3, i).text = f"{s['name']}\n{s['sub']}".strip()
                tbl.cell(2, i).paragraphs[0].runs[0].bold = True
        elif b["t"] == "witnesses":
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.add_run("WITNESSES:").bold = True
            for n in (1, 2):
                doc.add_paragraph(f"{n}. Signature: ______________________   Name: ______________________\n"
                                  f"    Address: ________________________________________________")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
def to_pdf(blocks, tpl_key, title, paper="legal", stamp_top_mm=100, stamp_line=None):
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from reportlab.pdfgen import canvas as rl_canvas
    from xml.sax.saxutils import escape
    w, h = PAPER.get(paper, PAPER["legal"])
    W, H = w * mm, h * mm
    initials = initials_line(tpl_key)

    class Numbered(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for st in self._pages:
                self.__dict__.update(st)
                self.setFont("Times-Roman", 8.5)
                self.setFillGray(0.3)
                self.drawString(32 * mm, 12 * mm, initials)
                self.drawRightString(W - 22 * mm, 12 * mm, f"Page {self._pageNumber} of {total}")
                super().showPage()
            super().save()
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=(W, H), leftMargin=32 * mm, rightMargin=22 * mm, topMargin=22 * mm,
                            bottomMargin=24 * mm, title=title, author="Sellandiamman Estates")
    base = ParagraphStyle("b", fontName="Times-Roman", fontSize=11.5, leading=16, alignment=TA_JUSTIFY, spaceAfter=6)
    center = ParagraphStyle("c", parent=base, alignment=TA_CENTER)
    ttl = ParagraphStyle("t", parent=center, fontName="Times-Bold", fontSize=15, leading=20, spaceAfter=4)
    clause = ParagraphStyle("cl", parent=base, leftIndent=9 * mm, firstLineIndent=-9 * mm)
    small = ParagraphStyle("s", parent=base, fontSize=9, leading=12)
    els = [Spacer(1, max(0, float(stamp_top_mm or 0)) * mm)]
    if stamp_line:
        els.append(Paragraph(escape(stamp_line), small))
    for b in blocks:
        tx = escape(b.get("text") or "")
        if b["t"] == "title":
            els.append(Paragraph(f"<u>{tx}</u>", ttl))
        elif b["t"] == "subtitle":
            els.append(Paragraph(f"({tx})", center))
        elif b["t"] == "heading":
            els += [Spacer(1, 6), Paragraph(f"<b><u>{tx}</u></b>", center)]
        elif b["t"] == "center":
            els.append(Paragraph(f"<b>{tx}</b>", center))
        elif b["t"] == "para":
            els.append(Paragraph(tx, base))
        elif b["t"] == "clause":
            m = re.match(r"^([A-Z][A-Z ,'&/-]{2,40}):\s*(.*)$", b["text"], re.S)
            body = f"<b>{escape(m.group(1))}:</b> {escape(m.group(2))}" if m else tx
            els.append(Paragraph(f"{b['n']}.&nbsp;&nbsp;&nbsp;{body}", clause))
        elif b["t"] == "signatures":
            cells = [[Paragraph("<br/><br/>______________________________", base) for _ in b["parties"]],
                     [Paragraph(f"<b>{escape(s['label'])}</b><br/>{escape(s['name'] or '')}<br/>{escape(s['sub'] or '')}", base) for s in b["parties"]]]
            t = Table(cells, colWidths=[(W - 54 * mm) / len(b["parties"])] * len(b["parties"]))
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
            els += [Spacer(1, 12), KeepTogether(t)]
        elif b["t"] == "witnesses":
            els += [Spacer(1, 14), KeepTogether([Paragraph("<b>WITNESSES:</b>", base)] + [
                Paragraph(f"{n}. Signature: ______________________ &nbsp; Name: ______________________<br/>"
                          f"&nbsp;&nbsp;&nbsp;&nbsp;Address: _______________________________________________", base) for n in (1, 2)])]
    doc.build(els, canvasmaker=Numbered)
    return buf.getvalue()
